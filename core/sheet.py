"""Download the community Thai translation (w3tu Google Sheets) and cache it."""
from __future__ import annotations

import gzip
import io
import json
import logging
import time
import urllib.request
from dataclasses import dataclass
from pathlib import Path

from .paths import assets_dir, cache_dir
from .progress import ProgressFn, noop, scaled

log = logging.getLogger(__name__)

EXPORT_URL = "https://docs.google.com/spreadsheets/d/{id}/export?format=xlsx"

# Later sources override earlier ones for the same string id.
SOURCES = [
    ("1lwoUZkMQcFwl_nlzV7UG9qxij6CMc5YuJOSsMuK3ZH8", ["v3", "v4"]),
    ("1YZvH6sd3XjPcLmz5nlRVkHd2Wm1psU1KthHtwSeG3t4", ["NextGen"]),
]

CACHE_NAME = "translations.json.gz"
CACHE_MAX_AGE = 24 * 3600
USER_AGENT = "ThaiW3Setup (+https://github.com/)"


@dataclass
class Translations:
    thai: dict[int, str]
    source: str
    fetched_at: float

    @property
    def age_text(self) -> str:
        return time.strftime("%d/%m/%Y %H:%M", time.localtime(self.fetched_at))


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(data: bytes, sheets: list[str]) -> dict[int, str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[int, str] = {}
    for name in sheets:
        if name not in wb.sheetnames:
            log.warning("sheet %s not found", name)
            continue
        for row in wb[name].iter_rows(min_row=3, max_col=5, values_only=True):
            if not row or row[0] is None or len(row) < 5:
                continue
            try:
                sid = int(str(row[0]).strip())
            except ValueError:
                continue
            thai = _cell_text(row[4])
            if thai.strip():
                out[sid] = thai
    wb.close()
    return out


def _download(url: str, progress: ProgressFn) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=120) as resp:
        total = int(resp.headers.get("Content-Length") or 0)
        buf = bytearray()
        while True:
            chunk = resp.read(256 * 1024)
            if not chunk:
                break
            buf += chunk
            if total:
                progress(len(buf) / total, f"กำลังดาวน์โหลดคำแปล {len(buf) // 1024:,} KB")
            else:
                progress(0.5, f"กำลังดาวน์โหลดคำแปล {len(buf) // 1024:,} KB")
    return bytes(buf)


def fetch_online(progress: ProgressFn = noop) -> dict[int, str]:
    merged: dict[int, str] = {}
    step = 1.0 / len(SOURCES)
    for i, (sheet_id, tabs) in enumerate(SOURCES):
        p = scaled(progress, i * step, (i + 0.7) * step)
        data = _download(EXPORT_URL.format(id=sheet_id), p)
        progress((i + 0.75) * step, "กำลังอ่านไฟล์คำแปล...")
        merged.update(parse_xlsx(data, tabs))
    return merged


def save_json(path: Path, thai: dict[int, str], fetched_at: float) -> None:
    payload = {"fetched_at": fetched_at, "strings": {str(k): v for k, v in thai.items()}}
    with gzip.open(path, "wt", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False)


def load_json(path: Path) -> tuple[dict[int, str], float]:
    with gzip.open(path, "rt", encoding="utf-8") as fh:
        payload = json.load(fh)
    return {int(k): v for k, v in payload["strings"].items()}, float(payload.get("fetched_at", 0))


def get_translations(force_download: bool = False, allow_online: bool = True,
                     progress: ProgressFn = noop) -> Translations:
    cache = cache_dir() / CACHE_NAME
    if cache.exists() and not force_download:
        try:
            thai, ts = load_json(cache)
            if time.time() - ts < CACHE_MAX_AGE or not allow_online:
                return Translations(thai, "คำแปลที่เก็บไว้ในเครื่อง", ts)
        except (OSError, ValueError, KeyError) as exc:
            log.warning("cache unreadable: %s", exc)

    if allow_online:
        try:
            thai = fetch_online(progress)
            ts = time.time()
            save_json(cache, thai, ts)
            return Translations(thai, "Google Sheet (ล่าสุด)", ts)
        except Exception as exc:  # network errors come in many types
            log.warning("online download failed: %s", exc)

    if cache.exists():
        thai, ts = load_json(cache)
        return Translations(thai, "คำแปลในเครื่อง (ออฟไลน์)", ts)
    bundled = assets_dir() / CACHE_NAME
    if bundled.exists():
        thai, ts = load_json(bundled)
        return Translations(thai, "คำแปลที่มากับโปรแกรม (ออฟไลน์)", ts)
    raise RuntimeError("ดาวน์โหลดคำแปลไม่สำเร็จ และไม่พบคำแปลสำรองในโปรแกรม")


if __name__ == "__main__":
    import argparse
    import sys

    ap = argparse.ArgumentParser(description="Download translations and write a gzip JSON snapshot")
    ap.add_argument("--export", type=Path, default=assets_dir() / CACHE_NAME)
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO)
    data = fetch_online(lambda f, m: print(f"\r{f * 100:5.1f}% {m}", end="", file=sys.stderr))
    args.export.parent.mkdir(parents=True, exist_ok=True)
    save_json(args.export, data, time.time())
    print(f"\nwrote {len(data):,} strings to {args.export}")
