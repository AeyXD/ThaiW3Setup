"""Text colours painted in a sheet become category <font color> tags."""
import io, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font

from core.custom import parse_custom_xlsx
from core.rich_color import (ALCHEMY, CATEGORY_COLORS, COMBAT, GENERAL, MUTATION, SIGNS, ColorWorkbook, category,
                             tag_runs)
from core.sheet import parse_text_xlsx, parse_xlsx

RED, BLUE, GREY = CATEGORY_COLORS[COMBAT], CATEGORY_COLORS[SIGNS], CATEGORY_COLORS[MUTATION]

# colours Google Sheets offers map to the category they look like; black, white and dark text stay normal
assert [category(c) for c in ("FFFF0000", "FFE06666", "FF0000FF", "FF6FA8DC", "FF00FF00", "FF93C47D",
                              "FFFFFF00", "FFFF9900", "FF999999", "FFB7B7B7", "FF9900FF", "FFFF00FF")] == \
    [COMBAT, COMBAT, SIGNS, SIGNS, ALCHEMY, ALCHEMY, GENERAL, GENERAL, MUTATION, MUTATION, SIGNS, COMBAT]
assert [category(c) for c in ("FF000000", "FF434343", "FFFFFFFF", "FFF3F3F3", None, "")] == [None] * 6

# runs: one tag per stretch of a colour, spaces outside, spaces between runs of one colour keep the tag whole
assert tag_runs([("A ", None), ("Adrenaline", "FFFF0000"), (" +1%", None)]) == f'A <font color="{RED}">Adrenaline</font> +1%'
assert tag_runs([(" Quen", "FF0000FF"), (" ", None), ("Sign ", "FF3C78D8")]) == f' <font color="{BLUE}">Quen Sign</font> '
assert tag_runs([("x", "FF000000")]) == "x"
assert tag_runs([('<font color="#CD7D03">Quen</font>', "FFFF0000")]) == '<font color="#CD7D03">Quen</font>'


def rich(*parts):
    return CellRichText(*[TextBlock(InlineFont(color=c), t) if c else t for t, c in parts])


# a custom sheet: part of a cell painted, a whole cell painted, black, grey, a theme colour and plain rows
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["title"])
ws.append(["ID", "TRANSLATE", "THAI"])
ws.append([1, rich(("\u0e04\u0e48\u0e32 ", None), ("Adrenaline", "FFFF0000"), (" +1%", None)), None])
ws.append([2, "Igni", "\u0e2d\u0e34\u0e01\u0e19\u0e35"])
ws["B4"].font = Font(color="FF4A86E8")
ws["C4"].font = Font(color="FF4A86E8")
ws.append([3, "plain", None])
ws["B5"].font = Font(color="FF000000")
ws.append([4, rich(("Mutation", "FF999999"), (" text", None)), None])
ws.append([5, "theme", None])
ws["B7"].font = Font(color=openpyxl.styles.colors.Color(theme=1))  # dark 1: black
ws.append([6, "7", None])
buf = io.BytesIO()
wb.save(buf)
data = buf.getvalue()

custom = parse_custom_xlsx(data)
assert custom.strings == {1: f'\u0e04\u0e48\u0e32 <font color="{RED}">Adrenaline</font> +1%',
                          2: f'<font color="{BLUE}">Igni</font>', 3: "plain",
                          4: f'<font color="{GREY}">Mutation</font> text', 5: "theme", 6: "7"}, custom.strings
assert custom.thai == {2: f'<font color="{BLUE}">\u0e2d\u0e34\u0e01\u0e19\u0e35</font>'}, custom.thai
assert custom.ids == [1, 2, 3, 4, 5, 6], custom.ids

rows = dict(ColorWorkbook(data).rows(0))
assert rows[3][0].text == "1" and rows[3][1].text == "\u0e04\u0e48\u0e32 Adrenaline +1%", rows[3]

# the w3tu layout: ids in A, Thai in E from row 3
main = openpyxl.Workbook()
mws = main.active
mws.title = "v4"
mws.append(["header painted"])
mws["A1"].font = Font(color="FFFF0000")
mws.append(["id", "", "", "en", "th"])
mws.append([10, None, None, "Muscle Memory", rich(("Adrenaline", "FF6AA84F"), (" \u0e40\u0e1e\u0e34\u0e48\u0e21", None))])
mws.append([11, None, None, "x", "  "])
buf = io.BytesIO()
main.save(buf)
assert parse_xlsx(buf.getvalue(), ["v4"]) == {10: f'<font color="{CATEGORY_COLORS[ALCHEMY]}">Adrenaline</font> \u0e40\u0e1e\u0e34\u0e48\u0e21'}

# the English-keyed sheet: whole cells were painted as review marks, so only part-cell colours count
text = openpyxl.Workbook()
text.active.title = "rules"
tws = text.create_sheet("content0")
for _ in range(4):
    tws.append(["head"])
tws.append(["Horse", "", "\u0e21\u0e49\u0e32"])
tws["C5"].font = Font(color="FF00FF00")
tws.append(["Adrenaline up", "", rich(("Adrenaline", "FFFF0000"), (" \u0e40\u0e1e\u0e34\u0e48\u0e21", None))])
buf = io.BytesIO()
text.save(buf)
assert parse_text_xlsx(buf.getvalue()) == {"Horse": "\u0e21\u0e49\u0e32",
                                           "Adrenaline up": f'<font color="{RED}">Adrenaline</font> \u0e40\u0e1e\u0e34\u0e48\u0e21'}

print("test_rich_color ok")
