"""Publish proper names (characters, places, quests, skills, monsters, items, Gwent cards, other)
as tabs of the community sheet.

Each tab uses the custom-sheet layout plus a THAI column: people type the Thai name there
and TRANSLATE becomes "English (Thai)" by formula, so enabling the tab in ThaiW3Setup
shows e.g. "Yennefer (Thai)". One row covers the string ids of a name that share a tab
(comma separated). Re-running keeps what people typed, also for rows someone moved to another tab.

Candidates are strings the translation keeps in English and that look like names. Each string id
goes to the tab of what the game uses it for (see game_index.py: its key name, journal section or
entity template), so "Triss Merigold" the journal entry and "Triss Merigold" the Gwent card get a
row each. Ids the game files do not explain fall back to a guess (place words, Thai place prefixes)
meant to be fixed by hand. Quests come from the quest sheet, skills from their string keys.
Setup: see export_untranslated.py.

    python devtools/export_names.py                       # refresh the tabs of the community sheet
    python devtools/export_names.py --dry-run             # report what would move, write nothing
    python devtools/export_names.py --sheet <test id> --from-sheet <community id>   # try it on a copy
    python devtools/export_names.py --no-upload --xlsx names.xlsx
"""
from __future__ import annotations

import argparse
import collections
import json
import os
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from core.custom import (COMMUNITY_ID, NAME_TABS, TAB_CHARACTERS, TAB_GWENT, TAB_ITEMS, TAB_MONSTERS, TAB_OTHER,
                         TAB_PLACES, TAB_QUESTS, TAB_SKILLS, Overrides, get_custom)
from core.game_detect import identify
from core.options import load_options
from core.sheet import get_translations
from core.text_builder import _load_merged, _thai_for
from core.w3strings import hash_key
from export_untranslated import _client
from game_index import GameIndex, build_index

QUEST_SHEET = "1X2VOm21x5ow_duSutaXV3AsCdz2YT1RVZZQoiPZHp14"
HEADER = ["ID", "ENGLISH", "THAI", "TRANSLATE", "NOTE"]
COLUMN_WIDTHS = [150, 300, 240, 360, 220]
UPDATED_TH = "\u0e2d\u0e31\u0e1b\u0e40\u0e14\u0e15"  # "updated"

PARTICLES = {"of", "the", "van", "de", "var", "aep", "von", "an", "la", "le", "du", "der", "and", "in", "on", "at"}
WORD = re.compile(r"^[A-Z][A-Za-z'\-.]*$")
PLACE_WORDS = {
    "village", "camp", "cave", "caves", "bridge", "inn", "estate", "castle", "palace", "harbor", "harbour",
    "crossroads", "mill", "farm", "farmstead", "tower", "ruins", "cemetery", "port", "gate", "square",
    "isle", "island", "islands", "fort", "fortress", "keep", "swamp", "forest", "woods", "lake", "grove",
    "hill", "hills", "mountain", "mountains", "pass", "road", "tavern", "vineyard", "bay", "bog", "marsh",
    "mine", "mines", "quarry", "outpost", "hut", "lighthouse", "temple", "sanctuary", "monastery",
    "district", "quarter", "garden", "gardens", "prison", "estate", "manor", "abbey", "shrine", "valley",
    "coast", "beach", "field", "fields", "meadow", "clearing", "dock", "docks", "wharf", "sewers", "well",
}
REGIONS = {"Velen", "Novigrad", "Oxenfurt", "Skellige", "Toussaint", "Beauclair", "Vizima", "Kaer Morhen",
           "White Orchard", "Nilfgaard", "Temeria", "Redania", "Kaedwen", "Aedirn", "Cintra", "Brokilon",
           "Ard Skellig", "Undvik", "An Skellig", "Spikeroog", "Faroe", "Hindarsfjall", "Kaer Trolde"}
# items, skills and alchemy that carry a proper noun ("Nilfgaardian Armor Set", "Glyph of Quen")
ITEM_WORDS = {
    "potion", "armor", "armour", "set", "glyph", "rune", "enhanced", "superior", "mastercrafted", "boosts",
    "intensity", "sword", "oil", "decoction", "mutagen", "trousers", "gauntlets", "boots", "blade", "bomb",
    "recipe", "diagram", "formula", "upgrade", "gear", "crossbow", "bolt", "bolts", "dust", "extract",
    "puffball", "essence", "saddle", "saddlebags", "blinders", "card", "cards", "location", "skill", "sign",
}
THAI_CHAR = re.compile("[\u0e00-\u0e7f]")
# the translation writes some places as "<Thai word> English", e.g. "city Novigrad"
THAI_PLACE_PREFIXES = {
    "\u0e40\u0e21\u0e37\u0e2d\u0e07", "\u0e14\u0e34\u0e19\u0e41\u0e14\u0e19",
    "\u0e2b\u0e21\u0e39\u0e48\u0e1a\u0e49\u0e32\u0e19", "\u0e40\u0e01\u0e32\u0e30",
    "\u0e1b\u0e23\u0e32\u0e2a\u0e32\u0e17", "\u0e16\u0e49\u0e33", "\u0e1b\u0e49\u0e2d\u0e21",
}

KEY_PREFIX_TABS = (
    ("gwint_", TAB_GWENT),
    ("item_name_", TAB_ITEMS), ("dlc_item_name_", TAB_ITEMS), ("effect_", TAB_ITEMS),
    ("glyphword", TAB_ITEMS),
    # map pins of the player and her horse, before the place pins
    ("map_location_player_", TAB_CHARACTERS), ("map_location_horse", TAB_CHARACTERS),
    ("location_name_", TAB_PLACES), ("map_location_", TAB_PLACES), ("cinematic_text_", TAB_PLACES),
    ("sign_name_", TAB_PLACES), ("poster_", TAB_PLACES),
    ("skill_", TAB_SKILLS), ("spell_power_", TAB_SKILLS),
    ("panel_hud_monstertype_", TAB_MONSTERS),
    ("leaderboard_name_", TAB_CHARACTERS),
    ("ea_", TAB_OTHER), ("input_", TAB_OTHER),  # achievements, keyboard keys
)
JOURNAL_TABS = (("bestiary", TAB_MONSTERS), ("characters", TAB_CHARACTERS), ("places", TAB_PLACES))
# quest folders like "monster_hunts" also hold the people who hand out the contract
MONSTER_ENTITY = re.compile(r"\\(monsters|animals)\\|\\[^\\]*_monster_[^\\]*$", re.I)
ANIMAL_ENTITY = re.compile(r"\\animals\\", re.I)
# where a key without a telling prefix was seen, checked in this order: monster entities, speaker
# names (voice tags, which also voice monsters), other entities, and plain mentions in scenes
SOURCE_TABS = (
    ("geralt_skills.xml", TAB_SKILLS),
    ("scene_voice_tags", TAB_CHARACTERS), ("attitudes.xml", TAB_CHARACTERS),
)
# credits list crew whose first names NPCs share, so they must not pull those NPCs along
WEAK_SOURCE_TABS = (
    (".w2scene", TAB_CHARACTERS), (".w2cutscene", TAB_CHARACTERS), (".w2comm", TAB_CHARACTERS),
    (".w2ent", TAB_CHARACTERS), ("credits", TAB_OTHER),
)
MINOR_JOURNAL = {"tutorial", "glossary", "storybook"}
# words in the NOTE column, the last resort for ids the game files do not explain at all; first match
NOTE_TABS = (
    ("\u0e21\u0e2d\u0e19\u0e2a\u0e40\u0e15\u0e2d\u0e23\u0e4c", TAB_MONSTERS),  # monster
    ("\u0e21\u0e34\u0e19\u0e34\u0e1a\u0e2d\u0e2a", TAB_MONSTERS),  # miniboss
    ("\u0e01\u0e32\u0e23\u0e4c\u0e14\u0e40\u0e01\u0e27\u0e19\u0e15\u0e4c", TAB_GWENT),  # gwent card
    ("\u0e2a\u0e16\u0e32\u0e19\u0e17\u0e35\u0e48", TAB_PLACES),  # place
    ("\u0e14\u0e32\u0e1a", TAB_ITEMS), ("\u0e40\u0e01\u0e23\u0e32\u0e30", TAB_ITEMS),  # sword, armor
    ("\u0e44\u0e2d\u0e40\u0e17\u0e21", TAB_ITEMS), ("\u0e27\u0e31\u0e15\u0e16\u0e38\u0e14\u0e34\u0e1a", TAB_ITEMS),  # item, ingredient
    ("\u0e04\u0e33\u0e28\u0e31\u0e1e\u0e17\u0e4c", TAB_OTHER), ("\u0e1a\u0e17\u0e1e\u0e39\u0e14", TAB_OTHER),  # term, line
    ("\u0e02\u0e49\u0e2d\u0e04\u0e27\u0e32\u0e21", TAB_OTHER), ("\u0e42\u0e23\u0e04", TAB_OTHER),  # UI text, disease
    ("\u0e22\u0e32", TAB_ITEMS),  # potion
)


def note_tab(note: str) -> str | None:
    for word, tab in NOTE_TABS:
        if word in note:
            return tab
    return None
# checked by hand: creatures the files only show as quest entities or speakers (most have a bestiary
# page under the plural, "Djinns", "Pixies", "Spriggans", or their own), and keyless items
HAND_TABS = {
    1208712: TAB_MONSTERS,  # Beast of Beauclair
    583032: TAB_MONSTERS,  # Djinn
    486014: TAB_MONSTERS,  # Drowned Dead
    1113321: TAB_MONSTERS,  # Incubus (the invisible rider)
    1075506: TAB_MONSTERS,  # Old Speartip (a cyclops)
    1163521: TAB_MONSTERS, 1186367: TAB_MONSTERS, 1177507: TAB_MONSTERS,  # Pixie
    345025: TAB_MONSTERS,  # Satyr
    1043714: TAB_MONSTERS,  # Silvan (Fugas's kind)
    1209957: TAB_MONSTERS,  # Spriggan
    1135402: TAB_ITEMS,  # Enervation, a glyphword amid the sign enchantments
    # tutorial titles shown with their subject's name
    1082705: TAB_CHARACTERS, 1084771: TAB_CHARACTERS,  # Ciri
    1196997: TAB_PLACES,  # Saint Lebioda's Footsteps
    # map pins of people
    388077: TAB_CHARACTERS, 1054862: TAB_CHARACTERS,  # Johnny ("The godling from Crookback Bog")
    1210126: TAB_CHARACTERS,  # Dye Merchant ("Seller of dyes.")
    # never referred to by the game files
    1041681: TAB_OTHER,  # Red Queen
    336331: TAB_OTHER, 169137: TAB_OTHER,  # Haalkatla, Hyphydria (cut content)
}
GWENT_RUN = 10
# the community highlights the THAI column of the characters tab
THAI_FILL = {TAB_CHARACTERS: {"red": 1, "green": 1}}
JOURNAL_NOTE = "journal"
# tabs filled from their own source, never emptied into others
FIXED_TABS = (TAB_QUESTS, TAB_SKILLS)


def evidence(sid: int, key: int | None, index: GameIndex) -> tuple[str | None, bool]:
    """(tab, strong) for what the game uses string ``sid`` (with key hash ``key``) for; tab is None
    when unknown. Weak answers come from an entity outside the monster folders, which holds people
    and monsters alike, or from a key merely mentioned in a scene."""
    name, sources = index.key_names.get(key, (None, [])) if key else (None, [])
    if name:
        for prefix, tab in KEY_PREFIX_TABS:
            if name.startswith(prefix):
                return tab, True
    sections = set(index.journal.get(sid, ()))
    for section, tab in JOURNAL_TABS:
        if section in sections:
            return tab, True
    entities = index.entities.get(sid)
    monster = bool(entities) and all(MONSTER_ENTITY.search(p) for p in entities)
    # a speaking animal (an owl that is a person) is a character; a speaking monster is not
    if monster and not any(ANIMAL_ENTITY.search(p) for p in entities):
        return TAB_MONSTERS, True
    # an entity whose name also keys monster loot is a monster spread over quest folders
    if entities and any("def_loot_monsters" in s for s in sources):
        return TAB_MONSTERS, True
    for marker, tab in SOURCE_TABS:
        if any(marker in s for s in sources):
            return tab, True
    if monster:
        return TAB_MONSTERS, True
    if entities:
        if any(MONSTER_ENTITY.search(p) and not ANIMAL_ENTITY.search(p) for p in entities):
            return TAB_MONSTERS, False
        return TAB_CHARACTERS, False
    for marker, tab in WEAK_SOURCE_TABS:
        if any(marker in s for s in sources):
            return tab, False
    if sections & MINOR_JOURNAL:
        return TAB_OTHER, True
    return None, False


def category(sid: int, key: int | None, index: GameIndex) -> str | None:
    return evidence(sid, key, index)[0]


def game_tabs(english, index: GameIndex) -> tuple[dict[int, str], set[int]]:
    """string id -> tab for every string the game files explain, and the ids explained only weakly."""
    out, weak = {}, set()
    for sid in english.strings:
        tab, strong = evidence(sid, english.keys.get(sid), index)
        if tab:
            out[sid] = tab
            if not strong:
                weak.add(sid)
    # Gwent cards sit in runs of name, quote, name, quote...; a card whose key the files never name
    # is still known by the cards around it
    for sid in english.strings:
        if sid in out:
            continue
        near = [out[n] for n in range(sid - GWENT_RUN, sid + GWENT_RUN + 1) if n in out and n not in weak]
        if near.count(TAB_GWENT) >= 2 and set(near) == {TAB_GWENT}:
            out[sid] = TAB_GWENT
    for sid, tab in HAND_TABS.items():
        if sid in english.strings:
            out[sid] = tab
            weak.discard(sid)
    return out, weak


def text_tabs(english, known: dict[int, str], weak: set[int]) -> dict[str, set[str]]:
    """English text -> tabs of the strongly explained strings with that text. Gwent cards are left
    out: they reuse the names of everything else."""
    out: dict[str, set[str]] = {}
    for sid, tab in known.items():
        if sid not in weak and tab != TAB_GWENT and sid in english.strings:
            out.setdefault(english.strings[sid].strip(), set()).add(tab)
    return out


def _title(tab: str) -> str:
    return f"ThaiW3Setup - {tab} ({UPDATED_TH} {time.strftime('%d/%m/%Y')})"


def looks_like_name(text: str) -> bool:
    words = text.split()
    if not 1 <= len(words) <= 4 or text.isupper() or text.endswith("-"):
        return False
    if any(re.search(r"[a-z][A-Z]|^[A-Z]{2}", w) for w in words):  # VSync, UI labels
        return False
    return bool(WORD.match(words[0])) and all(WORD.match(w) or w in PARTICLES for w in words)


def _slugs(name: str) -> set[str]:
    low = name.lower()
    plain = re.sub(r"[^a-z0-9 ]", "", low)
    return {low, low.replace(" ", "_"), low.replace(" ", ""), plain.replace(" ", "_"), plain.replace(" ", "")}


def _core_words(name: str) -> list[str]:
    parts = (p for w in name.split() if w not in PARTICLES for p in w.split("-"))
    return [p.strip("'.").lower() for p in parts if p.strip("'.")]


def lowercase_words(texts) -> set[str]:
    """Words used in lower case somewhere in the game text, i.e. ordinary words, not names."""
    out: set[str] = set()
    for t in texts:
        out.update(w.strip("'") for w in re.findall(r"[a-z][a-z']*", t))
    return out


def classify(name: str, keys: set[int], thai_place: bool, lower: set[str]) -> str | None:
    slugs = _slugs(name)
    if any(w in ITEM_WORDS for w in _core_words(name)):
        return None
    if (thai_place or name in REGIONS or any(hash_key("map_location_" + s) in keys for s in slugs)
            or name.split()[-1].lower() in PLACE_WORDS):
        return TAB_PLACES
    if any(hash_key(s) in keys for s in slugs):
        return TAB_CHARACTERS
    if any(w not in lower for w in _core_words(name)):
        return TAB_CHARACTERS
    return None


def skill_keys() -> list[str]:
    """String keys of the character skill tree, perks and mutations (not the basic attacks)."""
    keys = [f"skill_name_{tree}_s{n}" for tree in ("sword", "magic", "alchemy") for n in range(1, 31)]
    keys += [f"skill_name_perk_{n}" for n in range(1, 31)]
    keys += [f"skill_name_mutation_{n}" for n in range(1, 21)]
    return keys


def collect(english, thai: dict[int, str], by_text: dict[str, str] | None,
            quests: dict[int, str], known: dict[int, str] | None = None
            ) -> tuple[dict[str, dict[str, list[int]]], dict[str, str]]:
    """tab -> English name -> string ids, and English quest name -> current Thai (prefill).
    ``known`` is game_tabs(): the tab of each string id the game files explain."""
    known = known or {}
    none = Overrides()
    tabs: dict[str, dict[str, list[int]]] = {t: {} for t in NAME_TABS}
    prefill: dict[str, str] = {}
    for sid in sorted(quests):
        name = english.strings.get(sid, "").strip()
        if not name or len(name) > 80 or "<" in name:
            continue
        tabs[TAB_QUESTS].setdefault(name, []).append(sid)
        th = (_thai_for(sid, name, thai, by_text, none) or "").strip()
        if THAI_CHAR.search(th):
            prefill.setdefault(name, th)

    by_key = {k: sid for sid, k in english.keys.items()}
    skill_ids = [by_key[hash_key(k)] for k in skill_keys() if hash_key(k) in by_key]
    skill_names = {english.strings[sid].strip() for sid in skill_ids} - set(tabs[TAB_QUESTS]) - {""}
    for sid in skill_ids:
        th = (_thai_for(sid, english.strings[sid], thai, by_text, none) or "").strip()
        if THAI_CHAR.search(th):
            prefill.setdefault(english.strings[sid].strip(), th)

    found: dict[str, list[int]] = {}
    thai_place: set[str] = set()
    for sid, text in english.strings.items():
        name = text.strip()
        if name in skill_names:
            tabs[TAB_SKILLS].setdefault(name, []).append(sid)
            continue
        if name in tabs[TAB_QUESTS] or not looks_like_name(name):
            continue
        th = (_thai_for(sid, name, thai, by_text, none) or "").strip()
        if th == name:
            found.setdefault(name, []).append(sid)
        elif th.endswith(" " + name) and th[:-len(name)].strip() in THAI_PLACE_PREFIXES:
            found.setdefault(name, []).append(sid)
            thai_place.add(name)

    lower = lowercase_words(english.strings.values())
    for name, ids in found.items():
        unknown = [i for i in ids if i not in known]
        guess = None
        if unknown:
            keys = {english.keys[i] for i in ids if i in english.keys}
            guess = classify(name, keys, name in thai_place, lower)
        for sid in ids:
            tab = known.get(sid, guess)
            if tab:
                tabs[tab].setdefault(name, []).append(sid)
    for tab in tabs.values():
        for ids in tab.values():
            ids.sort()
    return tabs, prefill


def _existing(ws) -> dict[str, list]:
    """English -> [ids, thai, note] from a tab in our layout."""
    values = ws.get_all_values()
    header = next((i for i, row in enumerate(values[:10]) if "ENGLISH" in row and "THAI" in row), None)
    if header is None:
        return {}
    cols = {h: values[header].index(h) for h in HEADER if h in values[header]}
    out = {}
    for row in values[header + 1:]:
        cell = lambda h: row[cols[h]].strip() if h in cols and cols[h] < len(row) else ""
        if cell("ENGLISH"):
            ids = [int(float(p)) for p in cell("ID").split(",") if p.strip().replace(".", "", 1).isdigit()]
            out[cell("ENGLISH")] = [ids, cell("THAI"), cell("NOTE")]
    return out


def build_tabs(found: dict[str, dict[str, list[int]]], prefill: dict[str, str],
               existing: dict[str, dict[str, list]], known: dict[int, str] | None = None,
               weak: set[int] = frozenset(), same_text: dict[str, set[str]] | None = None
               ) -> dict[str, list[list]]:
    """Rows [ids, english, thai, note] per tab, one per name and tab. Each id goes to the tab
    ``known`` (game_tabs) gives it; ids the game files explain only weakly (``weak``) or not at all
    join the strongly explained ids of their name (in the rows, or failing that anywhere in the
    game per ``same_text``, text_tabs()) when those all went to one tab, and otherwise
    take their weak tab, the tab their NOTE names, or stay where they already are.
    Quests and skills keep their tab. A name whose ids end up in several tabs gets a row in each,
    all with the Thai and note typed so far; rows people filled in are kept even when the name
    is no longer found."""
    known = known or {}
    placed: dict[int, str] = {}
    placed_name: dict[str, str] = {}
    texts: dict[tuple[str, str], tuple[str, str]] = {}
    for tab in NAME_TABS:
        for name, (ids, thai, note) in existing.get(tab, {}).items():
            placed_name.setdefault(name, tab)
            texts[(tab, name)] = (thai, note)
            for sid in ids:
                placed.setdefault(sid, tab)
    by_name = {}
    for (tab, name), text in texts.items():
        if name not in by_name or (any(text) and not any(by_name[name])):
            by_name[name] = text
    # weakly or not explained ids follow the strongly explained ids of their name when those agree
    sibling_tabs: dict[str, set[str]] = {}
    for source in (found, {t: {n: v[0] for n, v in rs.items()} for t, rs in existing.items()}):
        for tab, names in source.items():
            if tab in FIXED_TABS:
                continue
            for name, ids in names.items():
                sibling_tabs.setdefault(name, set()).update(
                    known[s] for s in ids if s in known and s not in weak)

    def home_of(sid: int, name: str, fallback: str) -> str:
        if sid in known and sid not in weak:
            return known[sid]
        # a card or tutorial page named after an NPC does not make the NPC a card or a page
        ignore = {TAB_GWENT, TAB_OTHER} if sid in weak else set()
        tabs = (sibling_tabs.get(name, set()) - ignore) or ((same_text or {}).get(name, set()) - ignore)
        if len(tabs) == 1:
            return next(iter(tabs))
        return known.get(sid) or note_tab(by_name.get(name, ("", ""))[1]) or fallback

    rows: dict[str, dict[str, set[int]]] = {t: {} for t in NAME_TABS}
    done: set[str] = set()
    for tab in NAME_TABS:
        for name, ids in found.get(tab, {}).items():
            for sid in ids:
                if tab in FIXED_TABS:
                    home = tab
                else:
                    home = home_of(sid, name, placed.get(sid) or placed_name.get(name, tab))
                rows[home].setdefault(name, set()).add(sid)
            done.add(name)
    for tab in NAME_TABS:
        for name, (ids, thai, note) in existing.get(tab, {}).items():
            if name in done or not (thai or note) or not ids:
                continue
            for sid in ids:
                home = tab if tab in FIXED_TABS else home_of(sid, name, tab)
                rows[home].setdefault(name, set()).add(sid)

    out = {}
    for tab in NAME_TABS:
        out[tab] = []
        for name in sorted(rows[tab], key=str.lower):
            thai, note = texts.get((tab, name)) or by_name.get(name) or (prefill.get(name, ""), "")
            ids = sorted(rows[tab][name])
            out[tab].append([ids, name, thai, _note_for(note, tab, name, ids, rows, known)])
    return out


def _note_for(note: str, tab: str, name: str, ids: list[int],
              rows: dict[str, dict[str, set[int]]], known: dict[int, str]) -> str:
    """``note`` without the parts naming another tab that has its own row of ``name`` (they describe
    that row), nor "Gwent card" on a row none of whose ids the game uses on a card."""
    def wrong(part: str) -> bool:
        other = note_tab(part)
        if other in (None, tab):
            return False
        if name in rows[other]:
            return True
        return other == TAB_GWENT and not any(known.get(s) == TAB_GWENT for s in ids)
    parts = [p.strip() for p in note.split(",")]
    keep = [p for p in parts if not wrong(p)]
    return ", ".join(keep) if len(keep) < len(parts) else note


def journal_ids(index: GameIndex) -> set[int]:
    """Ids titling a page in the journal's characters section. Keys merely named by a journal do not
    count: quest entries name passers-by and monsters too."""
    return {sid for sid, sections in index.journal.items() if "characters" in sections}


def mark_journal(tabs: dict[str, list[list]], journal: set[int]) -> None:
    """NOTE the character rows the journal knows: on an empty NOTE, or over one naming another tab
    (the journal outranks it); free-form notes stay."""
    for row in tabs[TAB_CHARACTERS]:
        ids, note = row[0], row[3]
        if any(s in journal for s in ids) and (not note or note_tab(note)):
            row[3] = JOURNAL_NOTE


HAND_NOTES = Path(__file__).with_name("hand_notes.json")


def hand_notes(tabs: dict[str, list[list]], notes: dict[int, str]) -> None:
    """Put the game evidence checked by hand (id -> note) in the NOTE of the rows holding those ids."""
    for rows in tabs.values():
        for row in rows:
            found = [notes[s] for s in row[0] if s in notes]
            if found:
                row[3] = found[0]


def _id_cell(ids: list[int]):
    return ids[0] if len(ids) == 1 else ", ".join(map(str, ids))


def _formula(r: int) -> str:
    return f'=IF(C{r}="","",B{r}&" ("&C{r}&")")'


def write_xlsx(path: str, tabs: dict[str, list[list]]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, rows in tabs.items():
        ws = wb.create_sheet(tab)
        ws.append([_title(tab)])
        ws.append(HEADER)
        for r, (ids, name, thai, note) in enumerate(rows, 3):
            ws.append([_id_cell(ids), name, thai, _formula(r), note])
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDE", COLUMN_WIDTHS):
            ws.column_dimensions[col].width = width / 7
        ws.freeze_panes = "A3"
    wb.save(path)


def _fill_requests(sid: int, tab: str, rows: int, row_count: int) -> list[dict]:
    """Clear the backgrounds left by the previous rows, then paint the THAI column of the tabs that have it."""
    requests = [{"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 2, "endRowIndex": max(row_count, rows + 2)},
                                "cell": {"userEnteredFormat": {}},
                                "fields": "userEnteredFormat.backgroundColor"}}]
    if tab in THAI_FILL and rows:
        requests.append({"repeatCell": {
            "range": {"sheetId": sid, "startRowIndex": 2, "endRowIndex": rows + 2, "startColumnIndex": 2, "endColumnIndex": 3},
            "cell": {"userEnteredFormat": {"backgroundColor": THAI_FILL[tab]}},
            "fields": "userEnteredFormat.backgroundColor"}})
    return requests


def _format(sh, ws, tab: str, rows: int) -> None:
    sid = ws.id
    requests = _fill_requests(sid, tab, rows, ws.row_count) + [
        {"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {"frozenRowCount": 2}},
                                   "fields": "gridProperties.frozenRowCount"}},
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 1, "endRowIndex": 2},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
                        "fields": "userEnteredFormat.textFormat.bold"}},
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 2, "endRowIndex": rows + 2,
                                  "startColumnIndex": 0, "endColumnIndex": 5},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat.wrapStrategy,userEnteredFormat.verticalAlignment"}},
    ]
    for i, width in enumerate(COLUMN_WIDTHS):
        requests.append({"updateDimensionProperties": {
            "range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
            "properties": {"pixelSize": width}, "fields": "pixelSize"}})
    meta = sh.fetch_sheet_metadata({"fields": "sheets(properties.sheetId,protectedRanges,basicFilter)"})
    mine = next(s for s in meta["sheets"] if s["properties"]["sheetId"] == sid)
    protected = bool(mine.get("protectedRanges"))
    if mine.get("basicFilter"):
        requests.append({"setBasicFilter": {"filter": {"range": {
            "sheetId": sid, "startRowIndex": 1, "endRowIndex": rows + 2, "startColumnIndex": 0, "endColumnIndex": len(HEADER)}}}})
    if not protected:
        # "type the Thai name in the THAI column only"
        note = ("\u0e43\u0e2a\u0e48\u0e0a\u0e37\u0e48\u0e2d\u0e20\u0e32\u0e29\u0e32\u0e44\u0e17\u0e22"
                "\u0e43\u0e19\u0e04\u0e2d\u0e25\u0e31\u0e21\u0e19\u0e4c THAI \u0e40\u0e17\u0e48\u0e32\u0e19\u0e31\u0e49\u0e19")
        for start, end in ((0, 2), (3, 4)):
            requests.append({"addProtectedRange": {"protectedRange": {
                "range": {"sheetId": sid, "startColumnIndex": start, "endColumnIndex": end},
                "description": note, "warningOnly": True}}})
    sh.batch_update({"requests": requests})


def report(existing: dict[str, dict[str, list]], tabs: dict[str, list[list]], known: dict[int, str]) -> None:
    before = {sid: tab for tab in NAME_TABS for ids, _thai, _note in existing.get(tab, {}).values() for sid in ids}
    moves: collections.Counter = collections.Counter()
    for tab, rows in tabs.items():
        filled = sum(1 for r in rows if r[2])
        ids = [sid for r in rows for sid in r[0]]
        guessed = 0 if tab in FIXED_TABS else sum(1 for sid in ids if sid not in known)
        print(f"{ascii(tab)}: {len(rows):,} names, {len(ids):,} ids ({filled:,} names with Thai, "
              f"{guessed:,} ids not explained by the game files)")
        for sid in ids:
            if before.get(sid, tab) != tab:
                moves[(before[sid], tab)] += 1
    for (old, new), n in moves.most_common():
        print(f"  moved {n:,} ids {ascii(old)} -> {ascii(new)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--game", default=None, help="game folder (default: the one saved in settings)")
    ap.add_argument("--sheet", default=COMMUNITY_ID, help="sheet id to write the tabs to")
    ap.add_argument("--from-sheet", help="sheet id to read the typed Thai and notes from (default: --sheet)")
    ap.add_argument("--xlsx", help="also save the tabs to this .xlsx file")
    ap.add_argument("--no-upload", action="store_true", help="skip Google Sheets (use with --xlsx)")
    ap.add_argument("--dry-run", action="store_true", help="read the sheet and report the changes without writing")
    args = ap.parse_args()

    game = identify(args.game or load_options().game_path)
    if not game.supported:
        sys.exit(f"unsupported game folder: {game.label}")
    english = _load_merged(game, "en", [])
    index = build_index(game.path, english)
    known, weak = game_tabs(english, index)
    tr = get_translations()
    found, prefill = collect(english, tr.thai, tr.by_text, get_custom(QUEST_SHEET), known)

    sh = None
    sheets = {}
    existing: dict[str, dict[str, list]] = {}
    if not args.no_upload:
        client = _client()
        sh = client.open_by_key(args.sheet)
        sheets = {w.title: w for w in sh.worksheets()}
        source = sheets
        if args.from_sheet and args.from_sheet != args.sheet:
            source = {w.title: w for w in client.open_by_key(args.from_sheet).worksheets()}
        existing = {tab: _existing(source[tab]) for tab in NAME_TABS if tab in source}

    tabs = build_tabs(found, prefill, existing, known, weak, text_tabs(english, known, weak))
    mark_journal(tabs, journal_ids(index))
    hand_notes(tabs, {int(k): v for k, v in json.loads(HAND_NOTES.read_text(encoding="utf-8")).items()})
    report(existing, tabs, known)
    if args.xlsx:
        write_xlsx(args.xlsx, tabs)
        print(f"wrote {args.xlsx}")
    if sh is None or args.dry_run:
        return

    for tab, rows in tabs.items():
        ws = sheets.get(tab) or sh.add_worksheet(tab, rows=len(rows) + 2, cols=len(HEADER))
        ws.clear()
        if ws.row_count < len(rows) + 2 or ws.col_count < len(HEADER):
            ws.resize(rows=max(ws.row_count, len(rows) + 2), cols=max(ws.col_count, len(HEADER)))
        values = [[_title(tab)], HEADER] + [[_id_cell(ids), name, thai, "", note] for ids, name, thai, note in rows]
        ws.update(values, "A1", value_input_option="RAW")
        if rows:
            ws.update([[_formula(r)] for r in range(3, len(rows) + 3)], f"D3:D{len(rows) + 2}",
                      value_input_option="USER_ENTERED")
        _format(sh, ws, tab, len(rows))
        print(f"{ascii(tab)} gid={ws.id}")
    print(sh.url)


if __name__ == "__main__":
    main()
