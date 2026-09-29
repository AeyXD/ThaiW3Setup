"""Print tabs of a downloaded sheet xlsx: non-empty cells of the first rows plus row counts."""
import sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
tabs = sys.argv[2:] or wb.sheetnames
for name in tabs:
    ws = wb[name]
    n = 0
    samples = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(c is not None for c in row):
            n += 1
        if i <= 7:
            samples.append((i, [(j, ascii(str(c))[:50]) for j, c in enumerate(row) if c is not None]))
    print("==", ascii(name), "non-empty rows", n)
    for i, cells in samples:
        print("  ", i, cells)
