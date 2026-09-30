"""Name tabs: per-tab custom sheets, settings migration, "English (Thai)" in two-language mode."""
import json, os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

tmp = tempfile.mkdtemp()
os.environ["APPDATA"] = tmp  # keep the settings test away from the real profile

from core import custom
from core.custom import (COMMUNITY_ID, NAME_TABS, TAB_CHARACTERS, TAB_QUESTS, TAB_SKILLS, UNTRANSLATED_TAB,
                         _cache_path, default_sheets, parse_custom_xlsx, sheet_key)
from core.options import load_options, settings_path
from core.text_builder import combine
from export_names import build_tabs, classify, looks_like_name, skill_keys, write_xlsx

YEN_TH = "\u0e40\u0e22\u0e19\u0e40\u0e19\u0e40\u0e1f\u0e2d\u0e23\u0e4c"

# the workbook the exporter writes: one tab per kind, several ids per row
path = os.path.join(tmp, "names.xlsx")
write_xlsx(path, {TAB_CHARACTERS: [[[1, 2], "Yennefer", YEN_TH, ""], [[3], "Ciri", "", ""]],
                  NAME_TABS[1]: [], TAB_QUESTS: [[[9], "The Last Wish", "", "note"]]})
import openpyxl
wb = openpyxl.load_workbook(path)
ws = wb[TAB_CHARACTERS]
assert ws["D3"].value == '=IF(C3="","",B3&" ("&C3&")")', ws["D3"].value
ws["D3"] = f"Yennefer ({YEN_TH})"  # what Google exports: the computed value
wb.save(path)
data = open(path, "rb").read()
title, strings = parse_custom_xlsx(data, TAB_CHARACTERS)
assert strings == {1: f"Yennefer ({YEN_TH})", 2: f"Yennefer ({YEN_TH})"}, strings
assert parse_custom_xlsx(data, TAB_QUESTS)[1] == {}
try:
    parse_custom_xlsx(data, "missing")
    raise AssertionError("missing tab accepted")
except ValueError:
    pass

# tabs of one sheet get their own keys and caches; the old untranslated cache file stays in use
keys = [sheet_key(s) for s in default_sheets()]
assert len(keys) == len(set(keys)), keys
assert _cache_path(COMMUNITY_ID, UNTRANSLATED_TAB) == _cache_path(COMMUNITY_ID)
assert len({_cache_path(COMMUNITY_ID, t) for t in NAME_TABS}) == len(NAME_TABS)

# settings from 0.3.0 (bare community id, no tab) gain the name tabs once, disabled
old = [dict(s) for s in default_sheets() if not s.get("tab") or s["tab"] == UNTRANSLATED_TAB]
for s in old:
    s.pop("tab", None)
    s.pop("gid", None)
settings_path().write_text(json.dumps({"custom_sheets": old,
                                       "known_default_sheets": [s["sheet_id"] for s in old]}), encoding="utf-8")
opts = load_options()
community = [s for s in opts.custom_sheets if s["sheet_id"] == COMMUNITY_ID]
assert [s["tab"] for s in community] == [UNTRANSLATED_TAB, *NAME_TABS], community
assert community[0]["enabled"] and not any(s["enabled"] for s in community[1:])
opts.custom_sheets = [s for s in opts.custom_sheets if s.get("tab") != TAB_QUESTS]
settings_path().write_text(json.dumps({"custom_sheets": opts.custom_sheets,
                                       "known_default_sheets": opts.known_default_sheets}), encoding="utf-8")
assert TAB_QUESTS not in [s.get("tab") for s in load_options().custom_sheets]  # removal sticks

# two-language mode does not repeat the English name
assert combine(f"Yennefer ({YEN_TH})", "Yennefer", True) == f"Yennefer ({YEN_TH})"
assert combine(YEN_TH, "Yennefer", True) == f"{YEN_TH}  [Yennefer]"

# re-running the exporter keeps typed Thai, notes and hand-moved rows
found = {TAB_CHARACTERS: {"Yennefer": [1], "Novigrad": [5]}, NAME_TABS[1]: {}, TAB_QUESTS: {}}
existing = {NAME_TABS[1]: {"Novigrad": [[5], "TH", ""]}, TAB_CHARACTERS: {"Gone": [[7], "", "keep"]}}
tabs = build_tabs(found, {}, existing)
assert tabs[NAME_TABS[1]] == [[[5], "Novigrad", "TH", ""]], tabs
assert tabs[TAB_CHARACTERS] == [[[7], "Gone", "", "keep"], [[1], "Yennefer", "", ""]], tabs

# skills move to their tab even from a tab someone put them in, keeping the Thai
tabs = build_tabs({TAB_SKILLS: {"Whirl": [4, 8]}}, {}, {TAB_CHARACTERS: {"Whirl": [[4], "TH", ""]}})
assert tabs[TAB_SKILLS] == [[[4, 8], "Whirl", "TH", ""]] and tabs[TAB_CHARACTERS] == [], tabs
assert "skill_name_sword_s1" in skill_keys() and "skill_name_mutation_1" in skill_keys()

assert looks_like_name("Olgierd von Everec") and not looks_like_name("LEAD QA") and not looks_like_name("VSync")
assert classify("Nilfgaardian Armor Set", set(), False, set()) is None
assert classify("Crane Isle", set(), False, set()) == NAME_TABS[1]
assert classify("Food", set(), False, {"food"}) is None
print("test_names ok")
