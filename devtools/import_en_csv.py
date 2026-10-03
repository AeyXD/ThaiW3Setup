"""Publish every English string of an en.csv dump as the "ThaiW3 - rework" Google Sheet.

The sheet uses the custom-sheet layout (title row, then ID / ENGLISH / TRANSLATE / NOTE headers)
in the tab REWORK_TAB, so ThaiW3Setup can read it as a custom sheet (hidden until released,
see HIDDEN_SHEETS in core/custom.py). OAuth setup is the same as devtools/export_untranslated.py.

    python devtools/import_en_csv.py --csv en.csv                  # create a new sheet
    python devtools/import_en_csv.py --csv en.csv --sheet <id>     # refresh it, keeping translations
    python devtools/import_en_csv.py --csv en.csv --no-upload --xlsx rework.xlsx
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from core.custom import REWORK_NAME, REWORK_TAB as TAB, parse_sheet_id
from devtools.export_untranslated import COLUMN_WIDTHS, HEADER, _client, _existing_rows, _format

ROW_START = re.compile(r"^(\d+)\|[0-9A-Fa-f]*\|[^|]*\|")
BATCH_ROWS = 10_000


def parse_csv(text: str) -> dict[int, str]:
    """id -> English from ``id|key(hex)|key(str)|text`` lines; lines that do not start a row
    continue the text of the previous one (strings with line breaks). Later ids win."""
    out: dict[int, str] = {}
    last = None
    for line in text.lstrip("\ufeff").replace("\r\n", "\n").split("\n"):
        if line.startswith(";"):
            continue
        m = ROW_START.match(line)
        if m:
            last = int(m.group(1))
            out[last] = line[m.end():]
        elif last is not None and line:
            out[last] += "\n" + line
    return out


def build_rows(english: dict[int, str], existing: dict[int, list[str]]) -> list[list]:
    """Every id of the csv, keeping TRANSLATE / NOTE already in the sheet; rows only in the sheet stay
    when someone filled them."""
    rows = {}
    for sid, (old_english, translate, note) in existing.items():
        if sid in english or translate.strip() or note.strip():
            rows[sid] = [sid, english.get(sid, old_english), translate, note]
    for sid, text in english.items():
        rows.setdefault(sid, [sid, text, "", ""])
    return [rows[sid] for sid in sorted(rows)]


def _title() -> str:
    return f"{REWORK_NAME} (\u0e2d\u0e31\u0e1b\u0e40\u0e14\u0e15 {time.strftime('%d/%m/%Y')})"  # "updated"


def write_xlsx(path: str, rows: list[list]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TAB
    ws.append([_title()])
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCD", COLUMN_WIDTHS):
        ws.column_dimensions[col].width = width / 7
    for row in ws.iter_rows(min_row=3, min_col=2, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--csv", required=True, help="en.csv dumped from the game's w3strings")
    ap.add_argument("--sheet", help="existing sheet id or URL to refresh")
    ap.add_argument("--share", choices=["reader", "commenter", "writer"], default="reader",
                    help="access for anyone with the link")
    ap.add_argument("--xlsx", help="also save the rows to this .xlsx file")
    ap.add_argument("--no-upload", action="store_true", help="skip Google Sheets (use with --xlsx)")
    args = ap.parse_args()

    sheet_id = parse_sheet_id(args.sheet) if args.sheet else None
    if args.sheet and not sheet_id:
        sys.exit(f"not a sheet id or URL: {args.sheet}")

    with open(args.csv, encoding="utf-8-sig") as fh:
        english = parse_csv(fh.read())
    print(f"{len(english):,} strings in {args.csv}")

    sh = ws = None
    existing: dict[int, list[str]] = {}
    if not args.no_upload:
        gc = _client()
        if sheet_id:
            sh = gc.open_by_key(sheet_id)
            ws = next((w for w in sh.worksheets() if w.title == TAB), None)
            if ws is None:
                ws = sh.add_worksheet(TAB, rows=len(english) + 2, cols=len(HEADER))
            existing = _existing_rows(ws)
        else:
            sh = gc.create(REWORK_NAME)
            ws = sh.sheet1
            ws.update_title(TAB)

    rows = build_rows(english, existing)
    kept = sum(1 for r in rows if r[2].strip())
    print(f"{len(rows):,} rows ({kept:,} already translated in the sheet)")

    if args.xlsx:
        write_xlsx(args.xlsx, rows)
        print(f"wrote {args.xlsx}")
    if sh is None:
        return

    values = [[_title()], HEADER] + rows
    ws.clear()
    if ws.row_count < len(values) or ws.col_count < len(HEADER):
        ws.resize(rows=max(ws.row_count, len(values)), cols=max(ws.col_count, len(HEADER)))
    for start in range(0, len(values), BATCH_ROWS):
        ws.update(values[start:start + BATCH_ROWS], f"A{start + 1}", value_input_option="RAW")
        print(f"\ruploaded {min(start + BATCH_ROWS, len(values)):,}/{len(values):,} rows", end="")
    print()
    _format(sh, ws, len(rows))
    sh.share(None, perm_type="anyone", role=args.share, notify=False, with_link=True)
    print(f"sheet id: {sh.id}\ngid: {ws.id}\n{sh.url}")


if __name__ == "__main__":
    main()
