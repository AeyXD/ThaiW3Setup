"""Publish the English strings that have no Thai translation yet as a Google Sheet.

The sheet uses the custom-sheet layout (title row, then ID / TRANSLATE headers), so
once people translate rows it can be added to ThaiW3Setup as a custom sheet.

Setup (once): create a Google Cloud project, enable the Google Sheets API and the
Google Drive API, create an OAuth client of type "Desktop app" and save its JSON as
%APPDATA%\\ThaiW3Setup\\google_client.json. The first run opens a browser to sign in.

    pip install -r requirements-build.txt
    python devtools/export_untranslated.py                  # create a new sheet
    python devtools/export_untranslated.py --sheet <id>     # refresh it, keeping translations
    python devtools/export_untranslated.py --no-upload --xlsx untranslated.xlsx
"""
from __future__ import annotations

import argparse
import os
import sys
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from core.custom import UNTRANSLATED_TAB as TAB, default_sheets, merged_overrides, parse_sheet_id, sheet_key
from core.game_detect import identify
from core.options import load_options
from core.paths import app_data_dir
from core.sheet import get_translations
from core.text_builder import untranslated

HEADER = ["ID", "ENGLISH", "TRANSLATE", "NOTE"]
UNTRANSLATED_TH = "\u0e02\u0e49\u0e2d\u0e04\u0e27\u0e32\u0e21\u0e17\u0e35\u0e48\u0e22\u0e31\u0e07\u0e44\u0e21\u0e48\u0e41\u0e1b\u0e25"  # "untranslated text"
SPREADSHEET_NAME = f"ThaiW3Setup - {UNTRANSLATED_TH}"
COLUMN_WIDTHS = [90, 480, 480, 220]


def _title() -> str:
    return f"{SPREADSHEET_NAME} (\u0e2d\u0e31\u0e1b\u0e40\u0e14\u0e15 {time.strftime('%d/%m/%Y')})"  # "updated"


def _existing_rows(ws) -> dict[int, list[str]]:
    """id -> [english, translate, note] from a sheet in our layout."""
    values = ws.get_all_values()
    header = next((i for i, row in enumerate(values[:10]) if "ID" in row and "TRANSLATE" in row), None)
    if header is None:
        return {}
    cols = [values[header].index(h) if h in values[header] else None for h in HEADER]
    out = {}
    for row in values[header + 1:]:
        cell = lambda c: row[c] if c is not None and c < len(row) else ""
        try:
            sid = int(float(cell(cols[0])))
        except ValueError:
            continue
        out[sid] = [cell(cols[1]), cell(cols[2]), cell(cols[3])]
    return out


def build_rows(pending: dict[int, str], existing: dict[int, list[str]]) -> list[list]:
    """Keep every row someone translated or annotated; add ids that are still untranslated."""
    rows = {}
    for sid, (english, translate, note) in existing.items():
        if translate.strip() or note.strip() or sid in pending:
            rows[sid] = [sid, pending.get(sid, english), translate, note]
    for sid, english in pending.items():
        rows.setdefault(sid, [sid, english, "", ""])
    return [rows[sid] for sid in sorted(rows)]


def write_xlsx(path: str, rows: list[list]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TAB
    ws.append([_title()])
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCD", COLUMN_WIDTHS):
        ws.column_dimensions[col].width = width / 7
    for row in ws.iter_rows(min_row=3, min_col=2, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    wb.save(path)


def _format(sh, ws, rows: int) -> None:
    sid = ws.id
    requests = [
        {"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {"frozenRowCount": 2}},
                                   "fields": "gridProperties.frozenRowCount"}},
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 1, "endRowIndex": 2},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}},
                        "fields": "userEnteredFormat.textFormat.bold"}},
        {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 2, "endRowIndex": rows + 2,
                                  "startColumnIndex": 1, "endColumnIndex": 4},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat.wrapStrategy,userEnteredFormat.verticalAlignment"}},
    ]
    for i, width in enumerate(COLUMN_WIDTHS):
        requests.append({"updateDimensionProperties": {
            "range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
            "properties": {"pixelSize": width}, "fields": "pixelSize"}})
    meta = sh.fetch_sheet_metadata({"fields": "sheets(properties.sheetId,protectedRanges)"})
    protected = any(s.get("protectedRanges") for s in meta["sheets"] if s["properties"]["sheetId"] == sid)
    if not protected:
        requests.append({"addProtectedRange": {"protectedRange": {
            "range": {"sheetId": sid, "startColumnIndex": 0, "endColumnIndex": 2},
            # "ID and source text; edit only TRANSLATE / NOTE"
            "description": "ID \u0e41\u0e25\u0e30\u0e02\u0e49\u0e2d\u0e04\u0e27\u0e32\u0e21\u0e15\u0e49\u0e19\u0e09\u0e1a\u0e31\u0e1a "
                           "\u0e41\u0e01\u0e49\u0e40\u0e09\u0e1e\u0e32\u0e30\u0e04\u0e2d\u0e25\u0e31\u0e21\u0e19\u0e4c TRANSLATE / NOTE",
            "warningOnly": True}}})
    sh.batch_update({"requests": requests})


def _client():
    import gspread

    folder = app_data_dir()
    secret = folder / "google_client.json"
    if not secret.exists():
        sys.exit(f"missing OAuth client file: {secret}\n(see the setup notes at the top of this script)")
    return gspread.oauth(credentials_filename=str(secret), authorized_user_filename=str(folder / "google_token.json"))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--game", default=None, help="game folder (default: the one saved in settings)")
    ap.add_argument("--sheet", help="existing sheet id or URL to refresh")
    ap.add_argument("--share", choices=["reader", "commenter", "writer"], default="commenter",
                    help="access for anyone with the link")
    ap.add_argument("--xlsx", help="also save the rows to this .xlsx file")
    ap.add_argument("--no-upload", action="store_true", help="skip Google Sheets (use with --xlsx)")
    args = ap.parse_args()

    sheet_id = parse_sheet_id(args.sheet) if args.sheet else None
    if args.sheet and not sheet_id:
        sys.exit(f"not a sheet id or URL: {args.sheet}")

    game = identify(args.game or load_options().game_path)
    if not game.supported:
        sys.exit(f"unsupported game folder: {game.label}")
    tr = get_translations()
    # the name tabs of the same sheet still count as translations
    skip = {sheet_id, f"{sheet_id}#{TAB}"}
    sheets = [s for s in default_sheets() if sheet_key(s) not in skip]
    pending = untranslated(game, tr.thai, tr.by_text, merged_overrides(sheets))
    print(f"{len(pending):,} untranslated strings")

    sh = ws = None
    existing: dict[int, list[str]] = {}
    if not args.no_upload:
        gc = _client()
        if sheet_id:
            sh = gc.open_by_key(sheet_id)
            ws = next((w for w in sh.worksheets() if w.title == TAB), sh.sheet1)
            existing = _existing_rows(ws)
        else:
            sh = gc.create(SPREADSHEET_NAME)
            ws = sh.sheet1
            ws.update_title(TAB)

    rows = build_rows(pending, existing)
    kept = sum(1 for r in rows if r[2].strip())
    print(f"{len(rows):,} rows ({kept:,} already translated in the sheet)")

    if args.xlsx:
        write_xlsx(args.xlsx, rows)
        print(f"wrote {args.xlsx}")
    if sh is None:
        return

    values = [[_title()], HEADER] + rows
    ws.clear()
    if ws.row_count < len(values) or ws.col_count < len(HEADER):
        ws.resize(rows=max(ws.row_count, len(values)), cols=max(ws.col_count, len(HEADER)))
    ws.update(values, "A1", value_input_option="RAW")
    _format(sh, ws, len(rows))
    sh.share(None, perm_type="anyone", role=args.share, notify=False, with_link=True)
    print(f"sheet id: {sh.id}\n{sh.url}")


if __name__ == "__main__":
    main()
